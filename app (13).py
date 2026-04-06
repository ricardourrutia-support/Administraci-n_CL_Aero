import streamlit as st
import pandas as pd
import numpy as np
from datetime import datetime, timedelta, date
import io
import re
import xlsxwriter

# ─────────────────────────────────────────────
# CONFIGURACIÓN
# ─────────────────────────────────────────────
st.set_page_config(page_title="Gestor Operativo Aeropuerto", layout="wide")

# ── Header visual ──
import os
HEADER_PATH = "header.png"
if os.path.exists(HEADER_PATH):
    st.image(HEADER_PATH, use_container_width=True)

st.title("✈️ Gestor Operativo de Turnos — Aeropuerto")
st.caption("V1.0 — Lectura unificada · Colaciones automáticas · Excel limpio")

# ─────────────────────────────────────────────
# CONSTANTES Y REGLAS DE NEGOCIO
# ─────────────────────────────────────────────
ROLES_MAP = {
    "agente": "Agente",
    "ejecutivo": "Agente",
    "vendedor": "Agente",
    "anfitrion": "Anfitrión",
    "anfitrión": "Anfitrión",
    "anfitriones": "Anfitrión",
    "coordinador": "Coordinador",
    "coordinadores": "Coordinador",
    "supervisor": "Supervisor",
    "supervisores": "Supervisor",
}

LUGARES_AGENTE = ["Por Asignar", "T1 AIRE", "T1 TIERRA", "T2 AIRE", "T2 TIERRA", "Ausente"]
LUGARES_ANFITRION = ["Por Asignar", "Zona Nac", "Zona Int", "Ausente"]
LUGARES_COORD_SUP = ["General", "Ausente"]

IGNORE_NAMES = [
    "cargo", "nombre", "turno", "fecha", "total", "suma", "horas",
    "agente", "coordinador", "anfitrion", "anfitrión", "supervisor",
    "ejecutivo", "vendedor", ".", "..", "none", ""
]

NON_WORKING = [
    "libre", "l", "x", "vacaciones", "licencia", "falla",
    "domingos libres", "festivo", "feriado", "nan", "descanso",
    "permiso", "fuero"
]


# ─────────────────────────────────────────────
# FUNCIONES DE PARSEO
# ─────────────────────────────────────────────
def detect_role_from_sheet_name(sheet_name: str) -> str | None:
    """Detecta el rol a partir del nombre de la hoja."""
    s = sheet_name.lower().strip()
    for key, role in ROLES_MAP.items():
        if key in s:
            return role
    return None


def parse_shift(raw: str) -> tuple[list[int], int | None]:
    """
    Parsea un string de turno y devuelve (lista_de_horas, hora_inicio).
    Maneja: '08:00 - 19:00', '09:00-20:00', '22:00 - 09:00', '00:00 - 9:30'
    """
    if pd.isna(raw):
        return [], None
    s = str(raw).lower().strip()
    if s == "" or any(nw in s for nw in NON_WORKING):
        return [], None

    # Limpiar
    s = s.replace("hrs", "").replace("horas", "").replace("de ", "").replace("a ", "")
    s = s.replace("–", "-").replace("to", "-").replace("diurno", "").replace("nocturno", "")

    match = re.search(r'(\d{1,2})(?:[:.]\d+)?\s*-\s*(\d{1,2})(?:[:.]\d+)?', s)
    if not match:
        return [], None

    try:
        start_h = int(match.group(1))
        end_h = int(match.group(2))
    except ValueError:
        return [], None

    if not (0 <= start_h <= 23 and 0 <= end_h <= 23):
        return [], None

    if start_h == end_h:
        return [], None

    if start_h < end_h:
        hours = list(range(start_h, end_h))
    else:
        # Turno nocturno: 22:00 - 09:00
        hours = list(range(start_h, 24)) + list(range(0, end_h))

    return hours, start_h


def is_valid_name(val) -> bool:
    """Verifica si un valor es un nombre válido de persona."""
    if pd.isna(val):
        return False
    s = str(val).strip()
    if len(s) < 3:
        return False
    if s.lower() in IGNORE_NAMES:
        return False
    if s.replace(".", "", 1).replace(",", "").replace(" ", "").isdigit():
        return False
    # Filtrar fórmulas Excel
    if s.startswith("="):
        return False
    return True


def find_date_row_and_map(ws_data: list[list], sheet_name: str) -> tuple[int, dict]:
    """
    Busca la fila con fechas y construye un mapa col_index -> date.
    Soporta:
      - Fechas como datetime
      - Números de día (Supervisores Marzo)
    Devuelve (data_start_row, {col_idx: date})
    """
    date_map = {}
    data_start_row = None

    for row_idx, row in enumerate(ws_data[:10]):
        date_count = 0
        num_count = 0
        for col_idx, val in enumerate(row):
            if col_idx == 0:
                continue
            if isinstance(val, datetime):
                date_count += 1
            elif isinstance(val, (int, float)) and not pd.isna(val):
                n = int(val)
                if 1 <= n <= 31:
                    num_count += 1

        # Fila con fechas datetime
        if date_count >= 3:
            for col_idx, val in enumerate(row):
                if col_idx == 0:
                    continue
                if isinstance(val, datetime):
                    date_map[col_idx] = val.date()
            data_start_row = row_idx + 1
            # Verificar si la siguiente fila tiene "Cargo"/"Nombre" → datos empiezan 1 más
            if data_start_row < len(ws_data):
                next_first = ws_data[data_start_row][0]
                if isinstance(next_first, str) and next_first.strip().lower() in IGNORE_NAMES:
                    data_start_row += 1
            break

        # Fila con números de día (Supervisores Marzo)
        if num_count >= 10:
            # Necesitamos inferir mes/año del nombre de la hoja
            month_map = {
                "enero": 1, "febrero": 2, "marzo": 3, "abril": 4,
                "mayo": 5, "junio": 6, "julio": 7, "agosto": 8,
                "septiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12
            }
            year = datetime.now().year
            month = None
            sn_lower = sheet_name.lower()
            for m_name, m_num in month_map.items():
                if m_name in sn_lower:
                    month = m_num
                    break
            if month is None:
                continue

            prev_day = 0
            for col_idx, val in enumerate(row):
                if col_idx == 0:
                    continue
                if isinstance(val, (int, float)) and not pd.isna(val):
                    d = int(val)
                    if 1 <= d <= 31:
                        # Detectar cambio de mes
                        actual_month = month
                        actual_year = year
                        if d < prev_day and prev_day > 20:
                            # Cambio de mes (ej: 28 → 1)
                            pass  # ya estamos en el mes correcto
                        elif d > 20 and month > 1:
                            # Días del mes anterior al inicio
                            actual_month = month - 1
                            if actual_month == 0:
                                actual_month = 12
                                actual_year = year - 1
                        try:
                            date_map[col_idx] = date(actual_year, actual_month, d)
                        except ValueError:
                            pass
                        prev_day = d

            data_start_row = row_idx + 1
            # Saltar fila de días de semana si existe
            if data_start_row < len(ws_data):
                next_first = ws_data[data_start_row][0]
                if isinstance(next_first, str) and next_first.strip().lower() in [
                    "supervisor", "nombre", "cargo"
                ] + list(IGNORE_NAMES):
                    data_start_row += 1
            break

    return data_start_row, date_map


def read_sheet(file, sheet_name: str, role: str, date_start: date, date_end: date) -> pd.DataFrame:
    """Lee una hoja del Excel y extrae los datos de turnos."""
    file.seek(0)

    from openpyxl import load_workbook
    wb = load_workbook(file, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        return pd.DataFrame()

    ws = wb[sheet_name]
    ws_data = []
    for row in ws.iter_rows(values_only=True):
        ws_data.append(list(row))
    wb.close()

    if not ws_data:
        return pd.DataFrame()

    data_start_row, date_map = find_date_row_and_map(ws_data, sheet_name)

    if data_start_row is None or not date_map:
        return pd.DataFrame()

    # Filtrar solo fechas en el rango expandido (1 día antes para efecto cenicienta)
    load_start = date_start - timedelta(days=1)
    filtered_dates = {
        col: d for col, d in date_map.items()
        if load_start <= d <= date_end
    }

    records = []
    for row_idx in range(data_start_row, len(ws_data)):
        row = ws_data[row_idx]
        if not row or len(row) == 0:
            continue

        name_val = row[0]
        if not is_valid_name(name_val):
            continue

        name = str(name_val).strip().title()

        for col_idx, d in filtered_dates.items():
            if col_idx >= len(row):
                shift_raw = ""
            else:
                shift_raw = row[col_idx]

            if pd.isna(shift_raw):
                shift_raw = ""

            records.append({
                "Nombre": name,
                "Rol": role,
                "Fecha": d,
                "Turno_Raw": str(shift_raw).strip(),
            })

    return pd.DataFrame(records)


def auto_detect_sheets(file) -> list[tuple[str, str]]:
    """Detecta automáticamente las hojas y su rol."""
    file.seek(0)
    xl = pd.ExcelFile(file)
    detected = []
    for sn in xl.sheet_names:
        role = detect_role_from_sheet_name(sn)
        if role:
            detected.append((sn, role))
    return detected


# ─────────────────────────────────────────────
# REGLAS DE TAREAS POR ROL
# ─────────────────────────────────────────────

# Opciones de dropdown por rol (para horas activas dentro del turno)
# T1 = tarea principal (no afecta dotación)
# T2 = cobertura misma especialidad, otro lugar (no afecta dotación)
# T3 = administrativa coordinador (sí afecta dotación)
# T4 = cobertura cruzada entre especialidades (sí afecta dotación)
# C  = colación (sí afecta dotación)
# HE = hora extra (fuera de turno)
TASK_OPTIONS = {
    "Agente":      ["1", "2", "4", "C"],
    "Anfitrión":   ["1", "2", "C"],
    "Coordinador": ["1", "2", "3", "C", "4"],
    "Supervisor":  [],  # Sin dropdown, siempre T1
}

# Opción para horas fuera de turno (HHEE)
HHEE_OPTION = ["", "HE"]


def compute_tasks_for_shift(start_h: int, rol: str) -> dict[int, str]:
    """
    Devuelve un diccionario {hora: tarea} con las tareas especiales
    (todo lo que NO sea T1) para un turno dado.
    Las horas que no aparezcan en el dict son T1 por defecto.

    Reglas:
      - Supervisor: todo T1, sin colación en el sistema
      - Coordinador: 3 horas off (1 C + 2 T3) según turno:
          * Turno 10: h10=T3, h14=C, h15=T3
          * Turno 05: h11=T3, h12=C, h13=T3
          * Turno 21: h05=T3, h06=C, h07=T3
          * Turno 13: h17=T3, h18=C, h19=T3 (genérico: +4/+5/+6)
      - Agente/Anfitrión: 1 hora de C según ingreso:
          * <=8: C a las 12
          * 9: C a las 13
          * 10: C a las 14
          * 11: C a las 15
          * 18-20: C a las 2
          * 21: C a las 3
          * 22: C a las 4
          * 23: C a las 5
          * Otro: C a ingreso+4
    """
    specials = {}

    if rol == "Supervisor":
        # Todo T1, nada especial
        return specials

    elif rol == "Coordinador":
        # 3 horas off: T3 + C + T3
        if start_h == 10:
            specials[10] = "3"   # T3 al ingreso
            specials[14] = "C"   # Colación
            specials[15] = "3"   # T3
        elif start_h == 5:
            specials[11] = "3"   # T3
            specials[12] = "C"   # Colación
            specials[13] = "3"   # T3
        elif start_h == 21:
            specials[5] = "3"    # T3
            specials[6] = "C"    # Colación
            specials[7] = "3"    # T3
        elif start_h == 13:
            specials[17] = "3"
            specials[18] = "C"
            specials[19] = "3"
        else:
            # Genérico: off empieza a ingreso+4
            off_start = (start_h + 4) % 24
            specials[off_start] = "3"
            specials[(off_start + 1) % 24] = "C"
            specials[(off_start + 2) % 24] = "3"

    elif rol in ("Agente", "Anfitrión"):
        # 1 hora de colación
        if start_h is None:
            return specials
        if 0 <= start_h <= 8:
            specials[12] = "C"
        elif start_h == 9:
            specials[13] = "C"
        elif start_h == 10:
            specials[14] = "C"
        elif start_h == 11:
            specials[15] = "C"
        elif 18 <= start_h <= 20:
            specials[2] = "C"
        elif start_h == 21:
            specials[3] = "C"
        elif start_h == 22:
            specials[4] = "C"
        elif start_h == 23:
            specials[5] = "C"
        else:
            specials[(start_h + 4) % 24] = "C"

    return specials


# ─────────────────────────────────────────────
# MOTOR LÓGICO
# ─────────────────────────────────────────────
def build_hourly_grid(df_raw: pd.DataFrame, date_start: date, date_end: date) -> pd.DataFrame:
    """
    Construye la grilla hora a hora para cada persona/fecha.
    Retorna DataFrame con columnas:
      Nombre, Rol, Fecha, Turno_Raw, Hora, Tarea, Lugar
    """
    rows = []

    for _, r in df_raw.iterrows():
        hours, start_h = parse_shift(r["Turno_Raw"])
        nombre = r["Nombre"]
        rol = r["Rol"]
        fecha = r["Fecha"]

        if not hours:
            # Día libre o no laborable
            rows.append({
                "Nombre": nombre,
                "Rol": rol,
                "Fecha": fecha,
                "Shift_Date": fecha,
                "Turno_Raw": r["Turno_Raw"],
                "Hora": -1,
                "Tarea": r["Turno_Raw"] if r["Turno_Raw"] else "Libre",
                "Lugar": "",
                "Start_H": -1,
            })
            continue

        # Obtener tareas especiales para este turno
        specials = compute_tasks_for_shift(start_h, rol)

        for h in hours:
            # Calcular fecha real (efecto cenicienta)
            if start_h >= 18 and h < 12:
                fecha_real = fecha + timedelta(days=1)
            else:
                fecha_real = fecha

            if isinstance(fecha_real, datetime):
                fecha_real = fecha_real.date()
            if isinstance(fecha, datetime):
                fecha = fecha.date() if hasattr(fecha, 'date') else fecha

            # Tarea: especial si existe, T1 por defecto
            tarea = specials.get(h, "1")

            # Lugar por defecto
            if rol in ("Coordinador", "Supervisor"):
                lugar = "General"
            else:
                lugar = "Por Asignar"

            rows.append({
                "Nombre": nombre,
                "Rol": rol,
                "Fecha": fecha_real if isinstance(fecha_real, date) else fecha,
                "Shift_Date": fecha,  # día original del turno (para referenciar Lugar)
                "Turno_Raw": r["Turno_Raw"],
                "Hora": h,
                "Tarea": tarea,
                "Lugar": lugar,
                "Start_H": start_h,
            })

    df_grid = pd.DataFrame(rows)

    if df_grid.empty:
        return df_grid

    # Filtrar solo el rango pedido
    df_grid["Fecha"] = pd.to_datetime(df_grid["Fecha"]).dt.date
    df_grid = df_grid[
        (df_grid["Fecha"] >= date_start) & (df_grid["Fecha"] <= date_end)
    ]

    return df_grid


# ─────────────────────────────────────────────
# GENERADOR EXCEL (LIMPIO, SIN FÓRMULAS CRUZADAS)
# ─────────────────────────────────────────────
def generate_excel(df_grid: pd.DataFrame, date_start: date, date_end: date) -> io.BytesIO:
    """
    Genera el Excel de salida con:
      - Hoja Plan_Operativo con valores directos (no fórmulas cruzadas)
      - Formato condicional por colores
      - Dropdowns de lugar donde corresponda
      - Contadores de dotación activa
    """
    out = io.BytesIO()
    wb = xlsxwriter.Workbook(out, {"in_memory": True})

    # ── Formatos ──
    fmt = {
        "header": wb.add_format({
            "bold": True, "border": 1, "bg_color": "#7145D6",
            "font_color": "white", "align": "center", "valign": "vcenter",
            "font_size": 10
        }),
        "date_bar": wb.add_format({
            "bold": True, "border": 1, "bg_color": "#F0EDFF",
            "align": "center", "font_size": 10
        }),
        "base": wb.add_format({
            "border": 1, "align": "center", "font_size": 9, "text_wrap": True
        }),
        "group_title": wb.add_format({
            "bold": True, "border": 1, "bg_color": "#EFEFEF",
            "align": "left", "indent": 1, "font_size": 10
        }),
        "counter_label": wb.add_format({
            "bold": True, "border": 1, "bg_color": "#E8E0F7",
            "align": "center", "font_size": 9
        }),
        "counter_val": wb.add_format({
            "bold": True, "border": 1, "bg_color": "#E8E0F7",
            "align": "center", "font_size": 11
        }),
        "sep": wb.add_format({"bg_color": "#2D2D2D", "border": 0}),
        "libre": wb.add_format({
            "border": 1, "align": "center", "font_size": 9,
            "font_color": "#999999", "italic": True
        }),
        "alert": wb.add_format({
            "bg_color": "#EA9999", "font_color": "#980000",
            "bold": True, "border": 1, "align": "center"
        }),
    }

    task_colors = {
        "1": wb.add_format({"bg_color": "#D9EAD3", "border": 1, "align": "center", "font_size": 9}),
        "C": wb.add_format({"bg_color": "#FFF2CC", "border": 1, "align": "center", "font_size": 9, "bold": True}),
        "2": wb.add_format({"bg_color": "#CFE2F3", "border": 1, "align": "center", "font_size": 9}),
        "3": wb.add_format({"bg_color": "#D9D2E9", "border": 1, "align": "center", "font_size": 9}),
        "4": wb.add_format({"bg_color": "#FCE4D6", "border": 1, "align": "center", "font_size": 9}),
    }

    # ── Datos ──
    dates = sorted(df_grid["Fecha"].unique())
    dates = [d for d in dates if date_start <= d <= date_end]

    # Personas ordenadas por rol
    role_order = {"Agente": 1, "Anfitrión": 2, "Coordinador": 3, "Supervisor": 4}
    people = (
        df_grid[["Nombre", "Rol"]]
        .drop_duplicates()
        .assign(order=lambda x: x["Rol"].map(role_order).fillna(9))
        .sort_values(["order", "Nombre"])
    )

    # ── Hoja de validaciones (oculta, simple) ──
    ws_val = wb.add_worksheet("_Validaciones")
    ws_val.write_column(0, 0, LUGARES_AGENTE)           # Col A: Lugares Agente
    ws_val.write_column(0, 1, LUGARES_ANFITRION)         # Col B: Lugares Anfitrión
    ws_val.write_column(0, 2, TASK_OPTIONS["Agente"])    # Col C: Tareas Agente
    ws_val.write_column(0, 3, TASK_OPTIONS["Anfitrión"]) # Col D: Tareas Anfitrión
    ws_val.write_column(0, 4, TASK_OPTIONS["Coordinador"])  # Col E: Tareas Coordinador
    ws_val.write_column(0, 5, HHEE_OPTION)               # Col F: HE para fuera de turno
    ws_val.write_column(0, 6, LUGARES_COORD_SUP)         # Col G: Lugares Coord/Sup
    ws_val.hide()

    lugar_ag_range = f"_Validaciones!$A$1:$A${len(LUGARES_AGENTE)}"
    lugar_anf_range = f"_Validaciones!$B$1:$B${len(LUGARES_ANFITRION)}"
    task_ag_range = f"_Validaciones!$C$1:$C${len(TASK_OPTIONS['Agente'])}"
    task_anf_range = f"_Validaciones!$D$1:$D${len(TASK_OPTIONS['Anfitrión'])}"
    task_coord_range = f"_Validaciones!$E$1:$E${len(TASK_OPTIONS['Coordinador'])}"
    hhee_range = f"_Validaciones!$F$1:$F${len(HHEE_OPTION)}"
    lugar_cs_range = f"_Validaciones!$G$1:$G${len(LUGARES_COORD_SUP)}"

    # ── Hoja principal ──
    ws = wb.add_worksheet("Plan_Operativo")
    ws.set_tab_color("#7145D6")

    # Filas de conteo (arriba)
    ROW_COUNTER_AG = 0
    ROW_COUNTER_ANF = 1
    ROW_COUNTER_CO = 2
    ROW_COUNTER_SU = 3
    ROW_HEADER = 5
    ROW_HOURS = 6
    ROW_DATA_START = 7

    ws.write(ROW_COUNTER_AG, 0, "DOT. Agentes", fmt["counter_label"])
    ws.write(ROW_COUNTER_ANF, 0, "DOT. Anfitriones", fmt["counter_label"])
    ws.write(ROW_COUNTER_CO, 0, "DOT. Coordinadores", fmt["counter_label"])
    ws.write(ROW_COUNTER_SU, 0, "DOT. Supervisores", fmt["counter_label"])

    ws.write(ROW_HEADER, 0, "Colaborador", fmt["header"])
    ws.write(ROW_HEADER, 1, "Rol", fmt["header"])
    ws.set_column(0, 0, 22)
    ws.set_column(1, 1, 13)

    # Freezar paneles
    ws.freeze_panes(ROW_DATA_START, 2)

    dias_es = {0: "Lun", 1: "Mar", 2: "Mié", 3: "Jue", 4: "Vie", 5: "Sáb", 6: "Dom"}

    # ── Construir columnas por fecha (headers solamente, sin separadores aún) ──
    col = 2
    date_col_map = {}  # fecha -> col_start (columna de "Turno")
    sep_cols = []       # columnas separadoras para pintar al final
    # Cada bloque de fecha: [sep][Turno][Lugar][ShiftLugar(oculta)][h0..h23] = 1+1+1+1+24 = 28

    fmt_hidden = wb.add_format({"font_size": 1, "font_color": "#FFFFFF", "border": 0})

    for d in dates:
        dt = datetime.combine(d, datetime.min.time())
        d_label = f"{dias_es[dt.weekday()]} {dt.strftime('%d/%m')}"

        # Registrar columna separadora (se pintará después)
        sep_cols.append(col)
        ws.set_column(col, col, 1.5)
        col += 1

        # Header de fecha: merge sobre Turno + Lugar + ShiftLugar + 24 horas = 27 columnas
        ws.merge_range(ROW_HEADER, col, ROW_HEADER, col + 26, d_label, fmt["date_bar"])
        ws.write(ROW_HOURS, col, "Turno", fmt["header"])
        ws.write(ROW_HOURS, col + 1, "Lugar", fmt["header"])
        ws.write(ROW_HOURS, col + 2, "SL", fmt["header"])  # ShiftLugar (oculta)
        ws.set_column(col, col, 12)
        ws.set_column(col + 1, col + 1, 12)
        ws.set_column(col + 2, col + 2, 0.5)  # Ocultar columna SL (ancho mínimo)

        date_col_map[d] = col

        # Horas 0-23
        for h in range(24):
            h_col = col + 3 + h  # +3 por Turno, Lugar, ShiftLugar
            ws.write(ROW_HOURS, h_col, h, fmt["header"])
            ws.set_column(h_col, h_col, 4)

        col += 27  # turno(1) + lugar(1) + shiftlugar(1) + 24 horas

    total_cols = col  # ancho total de la grilla

    # ── Escribir datos persona por persona ──
    row = ROW_DATA_START
    current_role = None
    group_title_rows = []  # filas de título para NO pintar separadores ahí

    for _, person in people.iterrows():
        nombre = person["Nombre"]
        rol = person["Rol"]

        # Título de grupo
        if rol != current_role:
            label = {
                "Agente": "▸ AGENTES / EJECUTIVOS",
                "Anfitrión": "▸ ANFITRIONES",
                "Coordinador": "▸ COORDINADORES",
                "Supervisor": "▸ SUPERVISORES",
            }.get(rol, f"▸ {rol.upper()}")
            ws.write(row, 0, label, fmt["group_title"])
            ws.write(row, 1, "", fmt["group_title"])
            for c in range(2, total_cols):
                ws.write(row, c, "", fmt["group_title"])
            group_title_rows.append(row)
            row += 1
            current_role = rol

        ws.write(row, 0, nombre, fmt["base"])
        ws.write(row, 1, rol, fmt["base"])

        for d in dates:
            c_start = date_col_map[d]
            subset = df_grid[
                (df_grid["Nombre"] == nombre) & (df_grid["Fecha"] == d)
            ]

            working_hours = subset[subset["Hora"] != -1]
            is_working = not working_hours.empty

            # Determinar Shift_Date para esta persona/fecha
            # Si hay horas de madrugada de un nocturno del día anterior,
            # el Shift_Date será el día anterior
            shift_date_for_this_block = d
            if is_working:
                shift_dates_in_block = working_hours["Shift_Date"].unique()
                # Si algún Shift_Date es del día anterior, es un nocturno
                for sd in shift_dates_in_block:
                    if isinstance(sd, datetime):
                        sd = sd.date()
                    if sd < d:
                        shift_date_for_this_block = sd
                        break

            # Escribir la columna ShiftLugar como fórmula que referencia
            # la columna Lugar del Shift_Date
            if shift_date_for_this_block != d and shift_date_for_this_block in date_col_map:
                # Nocturno: referenciar Lugar del día anterior
                ref_lugar_col = date_col_map[shift_date_for_this_block] + 1
                ref_lugar_letter = xlsxwriter.utility.xl_col_to_name(ref_lugar_col)
                ws.write_formula(row, c_start + 2,
                    f'={ref_lugar_letter}{row + 1}', fmt_hidden)
            else:
                # Normal: referenciar Lugar de este mismo día
                lugar_letter = xlsxwriter.utility.xl_col_to_name(c_start + 1)
                ws.write_formula(row, c_start + 2,
                    f'={lugar_letter}{row + 1}', fmt_hidden)

            if not is_working:
                # Día libre — dropdown HE en cada hora
                turno_raw = ""
                libre_rows = subset[subset["Hora"] == -1]
                if not libre_rows.empty:
                    tr = libre_rows.iloc[0]["Tarea"]
                    turno_raw = tr if tr != "Libre" else ""

                ws.write(row, c_start, turno_raw if turno_raw else "Libre", fmt["libre"])
                ws.write(row, c_start + 1, "", fmt["libre"])
                for h in range(24):
                    cell_col = c_start + 3 + h
                    ws.write(row, cell_col, "", fmt["libre"])
                    ws.data_validation(row, cell_col, row, cell_col, {
                        "validate": "list", "source": hhee_range, "show_error": False
                    })
            else:
                # Día de trabajo
                turno_raw = working_hours.iloc[0].get("Turno_Raw", "")
                lugar = working_hours.iloc[0].get("Lugar", "Por Asignar")

                ws.write(row, c_start, str(turno_raw), fmt["base"])
                ws.write(row, c_start + 1, str(lugar), fmt["base"])

                # Dropdown de lugar
                if rol == "Agente":
                    ws.data_validation(row, c_start + 1, row, c_start + 1, {
                        "validate": "list", "source": lugar_ag_range, "show_error": False
                    })
                elif rol == "Anfitrión":
                    ws.data_validation(row, c_start + 1, row, c_start + 1, {
                        "validate": "list", "source": lugar_anf_range, "show_error": False
                    })
                elif rol in ("Coordinador", "Supervisor"):
                    ws.data_validation(row, c_start + 1, row, c_start + 1, {
                        "validate": "list", "source": lugar_cs_range, "show_error": False
                    })

                # Determinar rango de dropdown de tareas según rol
                task_range_for_role = {
                    "Agente": task_ag_range,
                    "Anfitrión": task_anf_range,
                    "Coordinador": task_coord_range,
                }.get(rol, None)

                # Escribir tareas hora a hora
                for h in range(24):
                    cell_col = c_start + 3 + h  # +3 por Turno, Lugar, ShiftLugar
                    h_data = working_hours[working_hours["Hora"] == h]
                    if h_data.empty:
                        # Hora fuera del turno → dropdown HE
                        ws.write(row, cell_col, "", fmt["base"])
                        ws.data_validation(row, cell_col, row, cell_col, {
                            "validate": "list", "source": hhee_range, "show_error": False
                        })
                    else:
                        tarea = str(h_data.iloc[0]["Tarea"])
                        cell_fmt = task_colors.get(tarea, fmt["base"])
                        ws.write(row, cell_col, tarea, cell_fmt)

                        # Dropdown de tarea (excepto Supervisores)
                        if task_range_for_role:
                            ws.data_validation(row, cell_col, row, cell_col, {
                                "validate": "list", "source": task_range_for_role,
                                "show_error": False
                            })

        row += 1

    # ── HHEE (filas manuales vacías) ──
    ws.write(row, 0, "▸ HORAS EXTRA / COBERTURAS MANUALES", fmt["group_title"])
    ws.write(row, 1, "", fmt["group_title"])
    for c in range(2, total_cols):
        ws.write(row, c, "", fmt["group_title"])
    group_title_rows.append(row)
    row += 1

    hhee_labels = [
        "HHEE Agente 1", "HHEE Agente 2", "HHEE Agente 3",
        "HHEE Anfitrión 1", "HHEE Coord/Sup 1"
    ]
    for label in hhee_labels:
        ws.write(row, 0, label, fmt["base"])
        ws.write(row, 1, "HHEE", fmt["base"])
        for d in dates:
            c_start = date_col_map[d]
            ws.write(row, c_start, "", fmt["base"])
            ws.write(row, c_start + 1, "Por Asignar", fmt["base"])
            ws.write(row, c_start + 2, "", fmt_hidden)  # ShiftLugar vacío
            for h in range(24):
                ws.write(row, c_start + 3 + h, "", fmt["base"])
        row += 1

    last_data_row = row - 1

    # ── Pintar columnas separadoras (ahora que ya sabemos el rango de filas) ──
    for sc in sep_cols:
        for r in range(ROW_HOURS, last_data_row + 1):
            if r not in group_title_rows:
                ws.write(r, sc, "", fmt["sep"])

    # ── Contadores de dotación (fórmulas COUNTIFS) ──
    # Dotación activa = (T1 + T2) excluyendo Ausentes
    # Un "Ausente" se marca en la columna Lugar de cada bloque de fecha
    fmt_counter_alert = wb.add_format({
        "bold": True, "border": 1, "bg_color": "#FF0000",
        "font_color": "white", "align": "center", "font_size": 11
    })

    for d in dates:
        c_start = date_col_map[d]
        # La columna "ShiftLugar" (SL) resuelve qué Lugar mirar para nocturnos
        sl_col_letter = xlsxwriter.utility.xl_col_to_name(c_start + 2)

        for h in range(24):
            h_col = c_start + 3 + h  # +3 por Turno, Lugar, ShiftLugar
            h_col_letter = xlsxwriter.utility.xl_col_to_name(h_col)
            r_start = ROW_DATA_START + 1  # +1 por el título de grupo
            r_end = last_data_row + 1     # xlsxwriter es 0-indexed

            rol_col = xlsxwriter.utility.xl_col_to_name(1)  # columna B = Rol
            data_range_rol = f"${rol_col}${r_start + 1}:${rol_col}${r_end + 1}"
            data_range_task = f"{h_col_letter}{r_start + 1}:{h_col_letter}{r_end + 1}"
            data_range_sl = f"{sl_col_letter}{r_start + 1}:{sl_col_letter}{r_end + 1}"

            for counter_row, rol_name in [
                (ROW_COUNTER_AG, "Agente"),
                (ROW_COUNTER_ANF, "Anfitrión"),
                (ROW_COUNTER_CO, "Coordinador"),
                (ROW_COUNTER_SU, "Supervisor"),
            ]:
                # Sumar T1 + T2 excluyendo filas con ShiftLugar = "Ausente"
                formula = (
                    f'=COUNTIFS({data_range_rol},"{rol_name}",{data_range_task},"1",{data_range_sl},"<>Ausente")'
                    f'+COUNTIFS({data_range_rol},"{rol_name}",{data_range_task},"2",{data_range_sl},"<>Ausente")'
                )
                ws.write_formula(counter_row, h_col, formula, fmt["counter_val"])

    # ── Formato condicional en contadores: rojo cuando bajo mínimos ──
    # Agentes < 4, Anfitriones < 2, Coordinadores = 0
    for d in dates:
        c_start = date_col_map[d]
        hours_range_start = c_start + 3    # +3 por Turno, Lugar, ShiftLugar
        hours_range_end = c_start + 26
        col_start_letter = xlsxwriter.utility.xl_col_to_name(hours_range_start)
        col_end_letter = xlsxwriter.utility.xl_col_to_name(hours_range_end)

        # Agentes < 4 → rojo
        ag_range = f"{col_start_letter}{ROW_COUNTER_AG + 1}:{col_end_letter}{ROW_COUNTER_AG + 1}"
        ws.conditional_format(ag_range, {
            "type": "cell", "criteria": "less than", "value": 4,
            "format": fmt_counter_alert
        })
        # Anfitriones < 2 → rojo
        anf_range = f"{col_start_letter}{ROW_COUNTER_ANF + 1}:{col_end_letter}{ROW_COUNTER_ANF + 1}"
        ws.conditional_format(anf_range, {
            "type": "cell", "criteria": "less than", "value": 2,
            "format": fmt_counter_alert
        })
        # Coordinadores = 0 → rojo
        co_range = f"{col_start_letter}{ROW_COUNTER_CO + 1}:{col_end_letter}{ROW_COUNTER_CO + 1}"
        ws.conditional_format(co_range, {
            "type": "cell", "criteria": "equal to", "value": 0,
            "format": fmt_counter_alert
        })

    # ── Formato condicional global ──
    data_area = f"A{ROW_DATA_START + 1}:{xlsxwriter.utility.xl_col_to_name(total_cols - 1)}{last_data_row + 1}"

    # Colación (amarillo)
    ws.conditional_format(data_area, {
        "type": "cell", "criteria": "equal to", "value": '"C"',
        "format": wb.add_format({"bg_color": "#FFF2CC", "border": 1, "align": "center", "bold": True})
    })
    # Tarea 2 - Cobertura misma especialidad, otro lugar (celeste)
    ws.conditional_format(data_area, {
        "type": "cell", "criteria": "equal to", "value": '"2"',
        "format": wb.add_format({"bg_color": "#CFE2F3", "border": 1, "align": "center"})
    })
    # Tarea 3 - Administrativa coordinador (lavanda)
    ws.conditional_format(data_area, {
        "type": "cell", "criteria": "equal to", "value": '"3"',
        "format": wb.add_format({"bg_color": "#D9D2E9", "border": 1, "align": "center"})
    })
    # Tarea 4 - Cobertura cruzada entre especialidades (naranja suave)
    ws.conditional_format(data_area, {
        "type": "cell", "criteria": "equal to", "value": '"4"',
        "format": wb.add_format({"bg_color": "#FCE4D6", "border": 1, "align": "center"})
    })
    # HE - Hora Extra (azul)
    ws.conditional_format(data_area, {
        "type": "cell", "criteria": "equal to", "value": '"HE"',
        "format": wb.add_format({"bg_color": "#4472C4", "font_color": "white", "border": 1, "align": "center", "bold": True})
    })

    wb.close()
    return out


# ═════════════════════════════════════════════
# INTERFAZ STREAMLIT
# ═════════════════════════════════════════════

st.sidebar.header("1. Cargar Archivo de Turnos")
st.sidebar.caption("Un solo Excel con hojas por rol (Agentes, Anfitriones, Coordinadores, Supervisores)")

uploaded_file = st.sidebar.file_uploader(
    "Archivo de turnos (.xlsx)", type=["xlsx"], key="turnos"
)

st.sidebar.markdown("---")
st.sidebar.header("2. Período a generar")

# Defaults: 1-15 del mes actual
today = datetime.now()
default_start = today.replace(day=1).date()
default_end = today.replace(day=15).date()

date_range = st.sidebar.date_input(
    "Rango de fechas",
    value=(default_start, default_end),
    format="DD/MM/YYYY"
)

start_d = None
end_d = None
if isinstance(date_range, (list, tuple)) and len(date_range) == 2:
    start_d, end_d = date_range

# ── Detección automática de hojas ──
sheet_assignments = {}
if uploaded_file:
    detected = auto_detect_sheets(uploaded_file)

    if detected:
        st.sidebar.markdown("---")
        st.sidebar.header("3. Hojas a incluir")

        month_names = {
            1: "enero", 2: "febrero", 3: "marzo", 4: "abril",
            5: "mayo", 6: "junio", 7: "julio", 8: "agosto",
            9: "septiembre", 10: "octubre", 11: "noviembre", 12: "diciembre"
        }

        target_month = month_names.get(start_d.month, "") if start_d else ""
        target_month_cap = target_month.capitalize() if target_month else ""

        # Separar hojas del mes principal vs hojas auxiliares
        primary_sheets = []
        auxiliary_sheets = []
        for sn, role in detected:
            if target_month and target_month.lower() in sn.lower():
                primary_sheets.append((sn, role))
            else:
                auxiliary_sheets.append((sn, role))

        # ── Hojas principales (del mes seleccionado) ──
        if primary_sheets:
            st.sidebar.markdown(f"**📅 Hojas de {target_month_cap}** *(mes del rango)*")
            for sn, role in primary_sheets:
                checked = st.sidebar.checkbox(
                    f"✅ {sn} → **{role}**",
                    value=True,
                    key=f"sheet_{sn}"
                )
                if checked:
                    sheet_assignments[sn] = role
        else:
            st.sidebar.warning(
                f"No se encontraron hojas para **{target_month_cap}**. "
                f"Verifica que el nombre de las hojas incluya el mes."
            )

        # ── Hojas auxiliares (otros meses, para rescatar madrugadas) ──
        if auxiliary_sheets:
            st.sidebar.markdown("---")
            st.sidebar.markdown(
                "**🌙 Hojas de otros meses** *(rescate de madrugadas)*"
            )
            st.sidebar.caption(
                f"Si el día anterior al inicio del rango ({(start_d - timedelta(days=1)).strftime('%d/%m/%Y') if start_d else '?'}) "
                f"no está incluido en las hojas de {target_month_cap}, "
                f"activa la hoja del mes anterior para rescatar los turnos nocturnos "
                f"que cruzan la medianoche (ej: 22:00-09:00)."
            )
            for sn, role in auxiliary_sheets:
                checked = st.sidebar.checkbox(
                    f"🌙 {sn} → **{role}**",
                    value=False,
                    key=f"sheet_{sn}"
                )
                if checked:
                    sheet_assignments[sn] = role

    else:
        st.sidebar.warning("No se detectaron hojas con nombres de roles conocidos.")

# ── Botón de generación ──
st.sidebar.markdown("---")

if st.sidebar.button("🚀 Generar Plan Operativo", type="primary", use_container_width=True):
    if not uploaded_file:
        st.error("⚠️ Carga un archivo Excel primero.")
    elif not start_d or not end_d:
        st.error("⚠️ Selecciona un rango de fechas válido.")
    elif not sheet_assignments:
        st.error("⚠️ Selecciona al menos una hoja.")
    else:
        with st.spinner("Leyendo hojas de turnos..."):
            all_dfs = []
            for sn, role in sheet_assignments.items():
                df_sheet = read_sheet(uploaded_file, sn, role, start_d, end_d)
                if not df_sheet.empty:
                    all_dfs.append(df_sheet)
                    st.sidebar.success(f"✓ {sn}: {df_sheet['Nombre'].nunique()} personas")
                else:
                    st.sidebar.warning(f"⚠ {sn}: sin datos en rango")

            if not all_dfs:
                st.error("No se encontraron datos en el rango seleccionado.")
            else:
                df_all = pd.concat(all_dfs, ignore_index=True)

                with st.spinner("Construyendo grilla horaria..."):
                    df_grid = build_hourly_grid(df_all, start_d, end_d)

                if df_grid.empty:
                    st.error("La grilla quedó vacía. Revisa el rango de fechas.")
                else:
                    # Stats
                    n_people = df_grid["Nombre"].nunique()
                    n_days = len(df_grid["Fecha"].unique())
                    n_by_role = df_grid[["Nombre", "Rol"]].drop_duplicates()["Rol"].value_counts()

                    col1, col2, col3 = st.columns(3)
                    col1.metric("Personas", n_people)
                    col2.metric("Días", n_days)
                    col3.metric("Registros hora", len(df_grid[df_grid["Hora"] != -1]))

                    st.markdown("**Distribución por rol:**")
                    for rol, count in n_by_role.items():
                        st.write(f"  • {rol}: {count}")

                    # Preview
                    with st.expander("👁️ Vista previa de la grilla (primeras 50 filas)"):
                        preview = df_grid[df_grid["Hora"] != -1][
                            ["Nombre", "Rol", "Fecha", "Hora", "Tarea", "Lugar"]
                        ].head(50)
                        st.dataframe(preview, use_container_width=True)

                    with st.spinner("Generando Excel..."):
                        excel_bytes = generate_excel(df_grid, start_d, end_d)

                    st.success("✅ ¡Plan Operativo generado exitosamente!")

                    fname = f"Plan_Operativo_{start_d.strftime('%d%b')}_{end_d.strftime('%d%b')}.xlsx"
                    st.download_button(
                        "📥 Descargar Plan Operativo",
                        data=excel_bytes.getvalue(),
                        file_name=fname,
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.xml",
                        type="primary",
                        use_container_width=True,
                    )
